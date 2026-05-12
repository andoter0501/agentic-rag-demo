from __future__ import annotations

import json
import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal


@dataclass
class Document:
    content: str
    metadata: dict[str, Any] = field(default_factory=dict)
    source: str = ""


@dataclass
class ParsedChunk:
    content: str
    metadata: dict[str, Any] = field(default_factory=dict)
    chunk_id: str = ""


class DocumentParser(ABC):
    @abstractmethod
    def parse(self, path: str | Path) -> list[Document]:
        raise NotImplementedError

    @abstractmethod
    def extract_chunks(
        self,
        document: Document,
        max_chars: int = 280,
        overlap: int = 50,
    ) -> list[ParsedChunk]:
        raise NotImplementedError


class MarkdownParser(DocumentParser):
    def parse(self, path: str | Path) -> list[Document]:
        p = Path(path)
        content = p.read_text(encoding="utf-8")
        return [Document(content=content, metadata={"source": str(p)}, source=str(p))]

    def extract_chunks(
        self,
        document: Document,
        max_chars: int = 280,
        overlap: int = 50,
    ) -> list[ParsedChunk]:
        from .loader import _split_sections, Chunk

        sections = _split_sections(document.content)
        chunks: list[ParsedChunk] = []

        for title, content in sections:
            paragraphs = [p.strip() for p in content.split("\n") if p.strip()]
            merged: list[str] = []
            bucket = ""

            for para in paragraphs:
                candidate = f"{bucket}\n{para}" if bucket else para
                if len(candidate) > max_chars and bucket:
                    merged.append(bucket)
                    bucket = para
                else:
                    bucket = candidate

            if bucket:
                merged.append(bucket)

            for idx, item in enumerate(merged, start=1):
                chunks.append(
                    ParsedChunk(
                        content=item,
                        metadata={
                            "title": title,
                            "source": document.source,
                            "chunk_index": idx,
                        },
                        chunk_id=f"{title}-{idx}",
                    )
                )

        return chunks


class PDFParser(DocumentParser):
    def __init__(self, extract_images: bool = False) -> None:
        self.extract_images = extract_images

    def parse(self, path: str | Path) -> list[Document]:
        try:
            import fitz
        except ImportError:
            raise RuntimeError("请安装 PyMuPDF: pip install pymupdf")

        doc = fitz.open(str(path))
        documents: list[Document] = []

        for page_num, page in enumerate(doc):
            text = page.get_text()
            if text.strip():
                documents.append(
                    Document(
                        content=text,
                        metadata={
                            "source": str(path),
                            "page": page_num + 1,
                            "total_pages": len(doc),
                        },
                        source=str(path),
                    )
                )

        doc.close()
        return documents

    def extract_chunks(
        self,
        document: Document,
        max_chars: int = 280,
        overlap: int = 50,
    ) -> list[ParsedChunk]:
        paragraphs = [p.strip() for p in document.content.split("\n\n") if p.strip()]
        chunks: list[ParsedChunk] = []
        bucket = ""
        page = document.metadata.get("page", 1)

        for idx, para in enumerate(paragraphs, start=1):
            if len(bucket) + len(para) > max_chars and bucket:
                chunks.append(
                    ParsedChunk(
                        content=bucket,
                        metadata={
                            "source": document.source,
                            "page": page,
                            "chunk_index": idx,
                        },
                        chunk_id=f"page{page}-chunk{idx}",
                    )
                )
                bucket = (
                    para[-overlap:] if overlap > 0 and len(para) > overlap else para
                )
            else:
                bucket = f"{bucket}\n{para}" if bucket else para

        if bucket:
            chunks.append(
                ParsedChunk(
                    content=bucket,
                    metadata={
                        "source": document.source,
                        "page": page,
                        "chunk_index": len(chunks) + 1,
                    },
                    chunk_id=f"page{page}-chunk{len(chunks) + 1}",
                )
            )

        return chunks


class HTMLParser(DocumentParser):
    def __init__(self, preserve_formatting: bool = True) -> None:
        self.preserve_formatting = preserve_formatting

    def parse(self, path: str | Path) -> list[Document]:
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            raise RuntimeError("请安装 beautifulsoup4: pip install beautifulsoup4")

        p = Path(path)
        html_content = p.read_text(encoding="utf-8")
        soup = BeautifulSoup(html_content, "html.parser")

        for script in soup(["script", "style"]):
            script.decompose()

        text = soup.get_text(separator="\n", strip=True)

        headers = []
        for h in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6"]):
            headers.append(h.get_text(strip=True))

        return [
            Document(
                content=text,
                metadata={
                    "source": str(p),
                    "title": soup.title.string if soup.title else str(p.stem),
                    "headers": headers,
                },
                source=str(p),
            )
        ]

    def extract_chunks(
        self,
        document: Document,
        max_chars: int = 280,
        overlap: int = 50,
    ) -> list[ParsedChunk]:
        paragraphs = [p.strip() for p in document.content.split("\n") if p.strip()]
        chunks: list[ParsedChunk] = []
        bucket = ""

        for idx, para in enumerate(paragraphs, start=1):
            if len(bucket) + len(para) > max_chars and bucket:
                chunks.append(
                    ParsedChunk(
                        content=bucket,
                        metadata={
                            "source": document.source,
                            "chunk_index": idx,
                        },
                        chunk_id=f"chunk-{idx}",
                    )
                )
                bucket = para[-overlap:] if overlap > 0 else para
            else:
                bucket = f"{bucket}\n{para}" if bucket else para

        if bucket:
            chunks.append(
                ParsedChunk(
                    content=bucket,
                    metadata={
                        "source": document.source,
                        "chunk_index": len(chunks) + 1,
                    },
                    chunk_id=f"chunk-{len(chunks) + 1}",
                )
            )

        return chunks


class JSONParser(DocumentParser):
    def parse(self, path: str | Path) -> list[Document]:
        p = Path(path)
        with open(p, "r", encoding="utf-8") as f:
            data = json.load(f)

        documents: list[Document] = []

        if isinstance(data, list):
            for i, item in enumerate(data):
                documents.append(
                    Document(
                        content=json.dumps(item, ensure_ascii=False),
                        metadata={"source": str(p), "index": i},
                        source=str(p),
                    )
                )
        elif isinstance(data, dict):
            documents.append(
                Document(
                    content=json.dumps(data, ensure_ascii=False),
                    metadata={"source": str(p)},
                    source=str(p),
                )
            )

        return documents

    def extract_chunks(
        self,
        document: Document,
        max_chars: int = 280,
        overlap: int = 50,
    ) -> list[ParsedChunk]:
        try:
            data = json.loads(document.content)
        except json.JSONDecodeError:
            return [
                ParsedChunk(
                    content=document.content,
                    metadata=document.metadata,
                    chunk_id="json-chunk-1",
                )
            ]

        chunks: list[ParsedChunk] = []

        if isinstance(data, dict):
            for key, value in data.items():
                chunk_content = f"{key}: {json.dumps(value, ensure_ascii=False)}"
                chunks.append(
                    ParsedChunk(
                        content=chunk_content,
                        metadata={
                            "source": document.source,
                            "key": key,
                        },
                        chunk_id=f"json-{key}",
                    )
                )

        elif isinstance(data, list):
            for idx, item in enumerate(data):
                chunk_content = json.dumps(item, ensure_ascii=False)
                if len(chunk_content) > max_chars:
                    chunk_content = chunk_content[:max_chars] + "..."
                chunks.append(
                    ParsedChunk(
                        content=chunk_content,
                        metadata={
                            "source": document.source,
                            "index": idx,
                        },
                        chunk_id=f"json-item-{idx}",
                    )
                )

        return chunks


class TXTParser(DocumentParser):
    def parse(self, path: str | Path) -> list[Document]:
        p = Path(path)
        content = p.read_text(encoding="utf-8")
        return [Document(content=content, metadata={"source": str(p)}, source=str(p))]

    def extract_chunks(
        self,
        document: Document,
        max_chars: int = 280,
        overlap: int = 50,
    ) -> list[ParsedChunk]:
        lines = [l.strip() for l in document.content.split("\n") if l.strip()]
        chunks: list[ParsedChunk] = []
        bucket = ""

        for idx, line in enumerate(lines, start=1):
            if len(bucket) + len(line) > max_chars and bucket:
                chunks.append(
                    ParsedChunk(
                        content=bucket,
                        metadata={
                            "source": document.source,
                            "line_start": idx - len(bucket.split("\n")),
                            "chunk_index": len(chunks) + 1,
                        },
                        chunk_id=f"txt-chunk-{len(chunks) + 1}",
                    )
                )
                bucket = (
                    line[-overlap:] if overlap > 0 and len(line) > overlap else line
                )
            else:
                bucket = f"{bucket}\n{line}" if bucket else line

        if bucket:
            chunks.append(
                ParsedChunk(
                    content=bucket,
                    metadata={
                        "source": document.source,
                        "chunk_index": len(chunks) + 1,
                    },
                    chunk_id=f"txt-chunk-{len(chunks) + 1}",
                )
            )

        return chunks


class DocxParser(DocumentParser):
    def parse(self, path: str | Path) -> list[Document]:
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise RuntimeError("请安装 python-docx: pip install python-docx")

        doc = DocxDocument(str(path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        return [
            Document(
                content="\n".join(paragraphs),
                metadata={
                    "source": str(path),
                    "paragraphs": len(paragraphs),
                },
                source=str(path),
            )
        ]

    def extract_chunks(
        self,
        document: Document,
        max_chars: int = 280,
        overlap: int = 50,
    ) -> list[ParsedChunk]:
        paragraphs = [p.strip() for p in document.content.split("\n") if p.strip()]
        chunks: list[ParsedChunk] = []
        bucket = ""

        for idx, para in enumerate(paragraphs, start=1):
            if len(bucket) + len(para) > max_chars and bucket:
                chunks.append(
                    ParsedChunk(
                        content=bucket,
                        metadata={
                            "source": document.source,
                            "chunk_index": idx,
                        },
                        chunk_id=f"docx-chunk-{idx}",
                    )
                )
                bucket = para[-overlap:] if overlap > 0 else para
            else:
                bucket = f"{bucket}\n{para}" if bucket else para

        if bucket:
            chunks.append(
                ParsedChunk(
                    content=bucket,
                    metadata={
                        "source": document.source,
                        "chunk_index": len(chunks) + 1,
                    },
                    chunk_id=f"docx-chunk-{len(chunks) + 1}",
                )
            )

        return chunks


def get_parser_for_file(path: str | Path) -> DocumentParser:
    suffix = Path(path).suffix.lower()

    parser_map: dict[str, type[DocumentParser]] = {
        ".md": MarkdownParser,
        ".markdown": MarkdownParser,
        ".pdf": PDFParser,
        ".html": HTMLParser,
        ".htm": HTMLParser,
        ".json": JSONParser,
        ".txt": TXTParser,
        ".docx": DocxParser,
    }

    parser_class = parser_map.get(suffix)
    if parser_class:
        return parser_class()

    return TXTParser()


def parse_document(
    path: str | Path,
    max_chars: int = 280,
    overlap: int = 50,
) -> list[ParsedChunk]:
    parser = get_parser_for_file(path)
    documents = parser.parse(path)
    all_chunks: list[ParsedChunk] = []

    for doc in documents:
        chunks = parser.extract_chunks(doc, max_chars=max_chars, overlap=overlap)
        all_chunks.extend(chunks)

    return all_chunks


def parse_documents(
    paths: list[str | Path],
    max_chars: int = 280,
    overlap: int = 50,
) -> list[ParsedChunk]:
    all_chunks: list[ParsedChunk] = []

    for path in paths:
        chunks = parse_document(path, max_chars=max_chars, overlap=overlap)
        all_chunks.extend(chunks)

    return all_chunks
